import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parents[1]
DOC_DIR = BASE_DIR / "Predictive_Modeling" / "docs"
OUTPUT_PATH = DOC_DIR / "ANN_Data_Validation_and_Architecture_Plan.docx"

NAVY = "1F4E78"
PALE_BLUE = "D9EAF7"
PALE_GREY = "F2F2F2"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margin(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    tr_pr.append(table_header)


def style_table(table, widths=None) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margin(cell)
            if widths:
                cell.width = widths[col_index]
            if row_index == 0:
                set_cell_shading(cell, NAVY)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
    style_table(table, widths=widths)
    document.add_paragraph()


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def metadata_lookup(metadata: pd.DataFrame, item: str) -> str:
    return str(metadata.loc[metadata["Item"] == item, "Value"].iloc[0])


def main() -> None:
    metadata = pd.read_csv(DOC_DIR / "ann_data_metadata.csv", dtype=str)
    checks = pd.read_csv(DOC_DIR / "ann_data_validation_checks.csv", dtype=str)
    distribution = pd.read_csv(DOC_DIR / "ann_class_distribution.csv", dtype=str)
    weights = pd.read_csv(DOC_DIR / "ann_class_weights.csv", dtype=str)
    config = json.loads((DOC_DIR / "ann_architecture_config.json").read_text(encoding="utf-8"))

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = BLACK
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Normal"].paragraph_format.line_spacing = 1.08

    for name, size in (("Title", 21), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("ANN Data Validation and Architecture Plan")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Customer Churn Analysis  Stage 3 Predictive Modelling")
    run.bold = True
    run.font.size = Pt(11)
    subtitle.paragraph_format.space_after = Pt(15)

    document.add_heading("Purpose and Decision", level=1)
    document.add_paragraph(
        "This document records the completed ANN data-readiness checks and the approved model architecture for Stage 3. "
        "The validated training and testing files are ready for model implementation. The preprocessing pipeline is fitted "
        "only on training data, the target and identifiers are excluded from inputs, and class imbalance is addressed in the training plan."
    )

    document.add_heading("Team Responsibilities", level=1)
    add_table(
        document,
        ["Member", "Role", "Stage 3 Responsibility"],
        [
            ["Ranjit Mishra", "Project Manager", "Coordinate work, maintain Jira and Confluence, check evidence, integrate the report, and manage submission."],
            ["Pratima Kandel", "Data Engineer", "Validate model-ready data, prevent preprocessing leakage, document class balance, and hand over consistent train and test files."],
            ["Muhammad Fahad Nazir", "ANN Technical Lead", "Define the ANN architecture, train the model, predict churn, evaluate performance, and save technical evidence."],
            ["SahilKumar Pravinbhai Patel", "Business Analyst", "Connect clustering and ANN findings to churn factors, retention actions, limitations, and business value."],
        ],
        [Inches(1.35), Inches(1.35), Inches(4.65)],
    )

    document.add_heading("Validated Dataset", level=1)
    dataset_rows = [
        ["Source rows", metadata_lookup(metadata, "Source rows")],
        ["Source columns", metadata_lookup(metadata, "Source columns")],
        ["Training rows", metadata_lookup(metadata, "Training rows")],
        ["Testing rows", metadata_lookup(metadata, "Testing rows")],
        ["Encoded input features", metadata_lookup(metadata, "Input features after encoding")],
        ["Training churn rate", metadata_lookup(metadata, "Training churn rate")],
        ["Testing churn rate", metadata_lookup(metadata, "Testing churn rate")],
        ["Random state", metadata_lookup(metadata, "Random state")],
    ]
    add_table(document, ["Measure", "Verified Value"], dataset_rows, [Inches(3.8), Inches(3.55)])

    document.add_heading("Validation Results", level=1)
    validation_rows = checks[["Check", "Status", "Evidence"]].values.tolist()
    add_table(
        document,
        ["Check", "Status", "Evidence"],
        validation_rows,
        [Inches(1.45), Inches(0.75), Inches(5.15)],
    )

    document.add_heading("Leakage Controls", level=2)
    add_bullet(document, "The raw data is divided into training and testing sets before preprocessing is fitted.")
    add_bullet(document, "The scaler and encoder are fitted only on training data and then applied to testing data.")
    add_bullet(document, "Churn is excluded from all model inputs and is used only as the prediction target.")
    add_bullet(document, "No generated customer identifier is used as a model input.")
    add_bullet(document, "Feature names and order are checked to match across training and testing files.")

    document.add_heading("Class Distribution", level=1)
    add_table(
        document,
        ["Split", "Class", "Count", "Percentage"],
        distribution[["Split", "Class", "Count", "Percentage"]].values.tolist(),
        [Inches(2.0), Inches(1.65), Inches(1.2), Inches(1.5)],
    )
    document.add_paragraph(
        "Churn is the minority class, so accuracy alone may hide poor detection of customers who leave. "
        "The evaluation will also report balanced accuracy, precision, recall, F1-score, ROC-AUC, and a confusion matrix."
    )

    document.add_heading("Suggested Class Weights", level=2)
    add_table(
        document,
        ["Class Value", "Class Label", "Balanced Weight"],
        weights[["Class Value", "Class Label", "Suggested Balanced Weight"]].values.tolist(),
        [Inches(1.35), Inches(2.5), Inches(2.0)],
    )

    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("ANN Architecture", level=1)
    architecture_rows = [
        ["Input", f"{config['input_features']} features", "Validated customer attributes"],
        ["Dense 1", "32 neurons  ReLU", "Learn main non-linear relationships"],
        ["Dropout 1", "Rate 0.20", "Reduce overfitting"],
        ["Dense 2", "16 neurons  ReLU", "Learn a smaller combined representation"],
        ["Dropout 2", "Rate 0.20", "Add a second overfitting control"],
        ["Output", "1 neuron  Sigmoid", "Produce churn probability from 0 to 1"],
    ]
    add_table(
        document,
        ["Layer", "Configuration", "Purpose"],
        architecture_rows,
        [Inches(1.35), Inches(2.15), Inches(3.85)],
    )
    flow = document.add_paragraph()
    flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    flow_run = flow.add_run("16 inputs  ->  Dense 32  ->  Dropout  ->  Dense 16  ->  Dropout  ->  Sigmoid")
    flow_run.bold = True
    flow_run.font.size = Pt(10.5)

    document.add_heading("Training Configuration", level=1)
    training_rows = [
        ["Optimizer", "Adam, learning rate 0.001"],
        ["Loss", "Binary cross-entropy"],
        ["Maximum epochs", "100"],
        ["Batch size", "32"],
        ["Validation", "20% of the training set"],
        ["Early stopping", "Monitor validation loss, patience 10, restore best weights"],
        ["Initial threshold", "0.50"],
        ["Reproducibility", "Random seed 42; save configuration, pipeline, code, model, and metrics"],
    ]
    add_table(document, ["Setting", "Decision"], training_rows, [Inches(2.05), Inches(5.3)])

    document.add_heading("Measurable Success Targets", level=1)
    targets = config["evaluation"]["success_targets"]
    success_rows = [
        ["ROC-AUC", f"At least {targets['roc_auc_minimum']:.2f}", "Measures ranking quality across thresholds."],
        ["Churn recall", f"At least {targets['churn_recall_minimum']:.2f}", "Prioritises finding customers at risk of leaving."],
        ["Churn F1-score", f"At least {targets['churn_f1_minimum']:.2f}", "Balances churn precision and recall."],
        ["Accuracy", f"At least {targets['accuracy_minimum']:.2f}", "Reported with class-sensitive metrics, not alone."],
        ["Overfitting control", "No severe divergence", "Compare training and validation loss curves."],
    ]
    add_table(
        document,
        ["Measure", "Project Target", "Reason"],
        success_rows,
        [Inches(1.55), Inches(1.6), Inches(4.2)],
    )
    document.add_paragraph(
        "These values are project targets rather than promised results. If a target is missed, the team will report the result honestly and review "
        "the threshold, class weights, hidden-layer size, dropout, and early-stopping behaviour without repeatedly tuning against the testing set."
    )

    document.add_heading("Risks and Controls", level=1)
    risk_rows = [
        ["Class imbalance", "Use balanced class weights and class-sensitive metrics."],
        ["Overfitting", "Use dropout, validation monitoring, early stopping, and learning curves."],
        ["Test-set leakage", "Fit preprocessing on training data only and reserve testing data for final evaluation."],
        ["Identical source rows", "Retain and disclose because no true customer ID is supplied; treat as a reporting limitation."],
        ["Weak business link", "Connect every recommendation to cluster evidence, ANN evidence, or both."],
    ]
    add_table(document, ["Risk", "Control"], risk_rows, [Inches(2.05), Inches(5.3)])

    document.add_heading("Required Training Evidence", level=1)
    for item in (
        "Trained ANN model and reproducible training code.",
        "Prediction probabilities and predicted churn classes for the testing data.",
        "Metrics table with accuracy, balanced accuracy, precision, recall, F1-score, and ROC-AUC.",
        "Confusion matrix, ROC curve, and training and validation learning curves.",
        "Short comparison against the measurable success targets.",
        "Evidence handover for the final report and demonstration video.",
    ):
        add_bullet(document, item)

    document.add_heading("Handover Decision", level=1)
    document.add_paragraph(
        "KAN-58 is technically complete: the model-ready datasets, preprocessing pipeline, validation checks, feature list, class distribution, "
        "and class weights have been produced. KAN-18 is technically complete: the ANN layers, training controls, evaluation measures, risks, and "
        "measurable targets have been defined. The next dependent task is ANN implementation, training, and evaluation under KAN-59."
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Customer Churn Analysis  Stage 3 ANN Readiness")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(89, 89, 89)

    document.core_properties.title = "ANN Data Validation and Architecture Plan"
    document.core_properties.subject = "Customer Churn Analysis Stage 3"
    document.core_properties.author = "Customer Churn Analysis Project Team"
    document.save(OUTPUT_PATH)
    print(f"Created {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
