from pathlib import Path
import csv
import re

from docx import Document


BASE_DIR = Path(__file__).resolve().parents[1]

REQUIRED_FILES = [
    "README.md",
    "requirements.txt",
    "Data_Preparation/raw/Dataset_ATS_v2.csv",
    "Data_Preparation/processed/cleaned_dataset.csv",
    "Data_Preparation/processed/preprocessed_dataset.csv",
    "Data_Preparation/processed/training_set.csv",
    "Data_Preparation/processed/testing_set.csv",
    "Data_Preparation/processed/missing_values_summary.csv",
    "Data_Preparation/processed/data_preparation_metadata.csv",
    "Data_Preparation/docs/data_preparation_summary.md",
    "Data_Preparation/docs/Data_Preparation_Documentation.docx",
    "Clustering_Analysis/docs/optimal_clusters_summary.csv",
    "Clustering_Analysis/docs/cluster_profile_summary.csv",
    "Clustering_Analysis/docs/customers_with_cluster_labels.csv",
    "Clustering_Analysis/docs/clustering_analysis_summary.md",
    "Clustering_Analysis/docs/Clustering_Analysis_Documentation.docx",
    "Clustering_Analysis/models/kmeans_model.pkl",
    "Clustering_Analysis/visualisations/elbow_method.png",
    "Clustering_Analysis/visualisations/silhouette_scores.png",
    "Clustering_Analysis/visualisations/cluster_size.png",
    "Clustering_Analysis/visualisations/churn_rate_by_cluster.png",
    "Clustering_Analysis/visualisations/tenure_monthlycharges_clusters.png",
    "scripts/01_data_preparation.py",
    "scripts/02_clustering_analysis.py",
]


def count_rows(path: Path) -> int:
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        next(reader)
        return sum(1 for _ in reader)


def docx_stats(path: Path) -> dict[str, int]:
    doc = Document(path)
    text = "\n".join(
        [p.text for p in doc.paragraphs]
        + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    )
    empty_cells = sum(
        1
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
        if not cell.text.strip()
    )
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes),
        "empty_cells": empty_cells,
        "words": len(re.findall(r"\b\w+\b", text)),
    }


def main() -> None:
    missing = [file for file in REQUIRED_FILES if not (BASE_DIR / file).exists()]
    if missing:
        print("Missing required files:")
        for file in missing:
            print(f"- {file}")
        raise SystemExit(1)

    print("All required files are present.")
    print(f"Raw rows: {count_rows(BASE_DIR / 'Data_Preparation/raw/Dataset_ATS_v2.csv')}")
    print(f"Preprocessed rows: {count_rows(BASE_DIR / 'Data_Preparation/processed/preprocessed_dataset.csv')}")
    print(f"Training rows: {count_rows(BASE_DIR / 'Data_Preparation/processed/training_set.csv')}")
    print(f"Testing rows: {count_rows(BASE_DIR / 'Data_Preparation/processed/testing_set.csv')}")
    print(f"Cluster-labelled rows: {count_rows(BASE_DIR / 'Clustering_Analysis/docs/customers_with_cluster_labels.csv')}")

    for doc in [
        BASE_DIR / "Data_Preparation/docs/Data_Preparation_Documentation.docx",
        BASE_DIR / "Clustering_Analysis/docs/Clustering_Analysis_Documentation.docx",
    ]:
        print(doc.name, docx_stats(doc))


if __name__ == "__main__":
    main()
