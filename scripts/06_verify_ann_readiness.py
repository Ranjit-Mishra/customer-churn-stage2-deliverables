import json
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
from docx import Document


BASE_DIR = Path(__file__).resolve().parents[1]
PM_DIR = BASE_DIR / "Predictive_Modeling"
DATA_DIR = PM_DIR / "data"
PREPROCESSOR_DIR = PM_DIR / "preprocessing"
DOC_DIR = PM_DIR / "docs"


def require_file(path: Path) -> None:
    if not path.exists() or path.stat().st_size == 0:
        raise AssertionError(f"Missing or empty output: {path.relative_to(BASE_DIR)}")


def main() -> None:
    required_files = [
        DATA_DIR / "ann_training_set.csv",
        DATA_DIR / "ann_testing_set.csv",
        PREPROCESSOR_DIR / "ann_preprocessor.joblib",
        DOC_DIR / "ann_feature_names.csv",
        DOC_DIR / "ann_class_distribution.csv",
        DOC_DIR / "ann_class_weights.csv",
        DOC_DIR / "ann_data_validation_checks.csv",
        DOC_DIR / "ann_data_metadata.csv",
        DOC_DIR / "ann_data_validation_summary.md",
        DOC_DIR / "ann_architecture_config.json",
        DOC_DIR / "ann_architecture_and_training_plan.md",
        DOC_DIR / "ANN_Data_Validation_and_Architecture_Plan.docx",
    ]
    for path in required_files:
        require_file(path)

    train = pd.read_csv(DATA_DIR / "ann_training_set.csv")
    test = pd.read_csv(DATA_DIR / "ann_testing_set.csv")
    feature_rows = pd.read_csv(DOC_DIR / "ann_feature_names.csv")
    checks = pd.read_csv(DOC_DIR / "ann_data_validation_checks.csv")
    metadata = pd.read_csv(DOC_DIR / "ann_data_metadata.csv", dtype=str)
    config = json.loads((DOC_DIR / "ann_architecture_config.json").read_text(encoding="utf-8"))
    preprocessor = joblib.load(PREPROCESSOR_DIR / "ann_preprocessor.joblib")

    assert len(train) == 5634, f"Unexpected training rows: {len(train)}"
    assert len(test) == 1409, f"Unexpected testing rows: {len(test)}"
    assert train.columns.tolist() == test.columns.tolist(), "Train/test columns do not match"
    assert train.columns[-1] == "Churn", "Target must be the final column"
    assert set(train["Churn"].unique()) == {0, 1}, "Training target is not binary"
    assert set(test["Churn"].unique()) == {0, 1}, "Testing target is not binary"
    assert not train.isna().any().any(), "Training set contains missing values"
    assert not test.isna().any().any(), "Testing set contains missing values"
    assert np.isfinite(train.drop(columns=["Churn"]).to_numpy()).all(), "Training inputs are not finite"
    assert np.isfinite(test.drop(columns=["Churn"]).to_numpy()).all(), "Testing inputs are not finite"

    feature_names = feature_rows["Feature"].tolist()
    assert feature_names == train.columns[:-1].tolist(), "Feature list does not match model-ready data"
    assert len(feature_names) == 16, f"Unexpected feature count: {len(feature_names)}"
    assert "Churn" not in feature_names, "Target leakage found"
    assert "CustomerID" not in feature_names, "Identifier leakage found"
    assert preprocessor.get_feature_names_out().tolist() == feature_names, "Saved pipeline feature order differs"
    assert config["input_features"] == len(feature_names), "Architecture input count differs from data"
    assert not (checks["Status"] == "FAIL").any(), "Validation checks contain a failure"

    metadata_map = dict(zip(metadata["Item"], metadata["Value"]))
    assert int(metadata_map["Training rows"]) == len(train)
    assert int(metadata_map["Testing rows"]) == len(test)
    assert int(metadata_map["Input features after encoding"]) == len(feature_names)

    document = Document(DOC_DIR / "ANN_Data_Validation_and_Architecture_Plan.docx")
    document_text = "\n".join(p.text for p in document.paragraphs)
    for phrase in (
        "Purpose and Decision",
        "Team Responsibilities",
        "Validated Dataset",
        "Validation Results",
        "ANN Architecture",
        "Measurable Success Targets",
        "Handover Decision",
    ):
        assert phrase in document_text, f"Word document is missing: {phrase}"
    assert all(cell.text.strip() for table in document.tables for row in table.rows for cell in row.cells), (
        "Word document contains an empty table cell"
    )

    print("ANN readiness verification passed.")
    print(f"Required files checked: {len(required_files)}")
    print(f"Training rows: {len(train):,}")
    print(f"Testing rows: {len(test):,}")
    print(f"Model input features: {len(feature_names)}")
    print(f"Validation statuses: {checks['Status'].value_counts().to_dict()}")
    print(f"Word document tables: {len(document.tables)}")


if __name__ == "__main__":
    main()
